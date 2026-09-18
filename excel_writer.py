
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import OUTPUT_COLUMNS, OUTPUT_SHEET_NAME
from parser import EmployeeRecord


def _autofit(ws):
    for col_idx, cells in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value)) for c in cells if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 48)


def build_workbook(records: list[EmployeeRecord]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = OUTPUT_SHEET_NAME

    headers = list(OUTPUT_COLUMNS)
    # Append all daily columns needed for this month's calendar.
    max_days = max((len(r.status_tokens) for r in records), default=31)
    headers.extend(f"Day {i}" for i in range(1, max_days + 1))
    ws.append(headers)

    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        base = record.to_output_row()
        row = [base.get(col, "") for col in OUTPUT_COLUMNS]
        row.extend(base.get(f"Day {i}", "") for i in range(1, max_days + 1))
        ws.append(row)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for idx, cell in enumerate(row, start=1):
            cell.alignment = Alignment(
                horizontal="center" if idx in {1,3,4,5,6,8,9,10,11,12} or idx > len(OUTPUT_COLUMNS) else "left",
                vertical="center",
                wrap_text=True,
            )
    _autofit(ws)

    return_io = io.BytesIO()
    wb.save(return_io)
    return_io.seek(0)
    return return_io
